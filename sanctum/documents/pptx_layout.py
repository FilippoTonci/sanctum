"""Review-surface layout for .pptx: ``GET /review-sessions/<id>/layout``.

Builds the shared layout contract (``LayoutResponse`` in
:mod:`sanctum.api.schemas`) from a presentation: one page per slide,
items in paint order, geometry in points with a top-left origin.

Segment ids come from the same walker the Reader/Writer use
(:mod:`sanctum.documents.pptx_adapter`), so every ``segment_id`` in the
layout is a segment in the session and vice versa.

Additions to the contract (all optional, additive):

- ``textbox.anchor``: ``"top" | "middle" | "bottom"`` — vertical anchor
  of the text frame (``bodyPr@anchor``); titles are usually bottom- or
  middle-anchored and look wrong top-aligned.
- ``image.alt``: ``{"segment_id", "text"}`` — picture alt-text segment.
- ``page.notes``: ``[Paragraph]`` — speaker-notes paragraphs (same shape
  as textbox paragraphs). Notes have no slide geometry.
- ``unscanned[].page``: 0-based page index the entry belongs to, or
  ``null`` for document-level entries.

This is a best-effort visual approximation, not a renderer: theme
colours, gradients, rotation, line art and bullets are not reproduced.
Text stays real text so the review highlight machinery works on it.
"""

from __future__ import annotations

import base64
from io import BytesIO
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.dml import MSO_COLOR_TYPE, MSO_FILL
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import PP_ALIGN

from sanctum.documents.pptx_adapter import (
    A_NS,
    P_NS,
    Transform,
    alt_text,
    is_group,
    iter_slide_shapes,
    notes_text_frame,
    run_id,
)

EMU_PER_PT = 12700.0
DEFAULT_FONT_PT = 18.0
_BROWSER_IMAGE_TYPES = {
    "image/png",
    "image/jpeg",
    "image/gif",
    "image/bmp",
    "image/svg+xml",
    "image/webp",
}
_ALIGN = {
    PP_ALIGN.LEFT: "left",
    PP_ALIGN.CENTER: "center",
    PP_ALIGN.RIGHT: "right",
    PP_ALIGN.JUSTIFY: "justify",
    PP_ALIGN.DISTRIBUTE: "justify",
    PP_ALIGN.JUSTIFY_LOW: "justify",
    PP_ALIGN.THAI_DISTRIBUTE: "justify",
}
_ALGN_ATTR = {"l": "left", "ctr": "center", "r": "right", "just": "justify", "dist": "justify"}
_ANCHOR_ATTR = {"t": "top", "ctr": "middle", "b": "bottom"}
_TITLE_PH_TYPES = {"title", "ctrTitle"}
# Layout placeholder types inherit from a differently-typed master placeholder.
_MASTER_PH_FOR = {
    PP_PLACEHOLDER.CENTER_TITLE: PP_PLACEHOLDER.TITLE,
    PP_PLACEHOLDER.SUBTITLE: PP_PLACEHOLDER.BODY,
}
_BODY_PH_TYPES = {"body", "subTitle", "obj", None}
# Neutral table cell fills when the cell has no explicit solid colour
# (table styles are theme-driven and not resolved here).
_TABLE_HEADER_FILL = "#D9E2F3"
_TABLE_BODY_FILL = "#F2F2F2"
_UNRENDERED_FILL = "#E7E6E6"

_GRAPHIC_DATA_KINDS = {
    "http://schemas.openxmlformats.org/drawingml/2006/chart": "Chart",
    "http://schemas.openxmlformats.org/drawingml/2006/diagram": "SmartArt diagram",
    "http://schemas.openxmlformats.org/presentationml/2006/ole": "Embedded object (OLE)",
}


def _pt(emu: float) -> float:
    return round(emu / EMU_PER_PT, 2)


def _box(shape: Any, xf: Transform, page_w: float, page_h: float) -> dict[str, float]:
    left, top, width, height = shape.left, shape.top, shape.width, shape.height
    if left is None or top is None or width is None or height is None:
        # Geometry not resolvable (e.g. placeholder whose layout lost its
        # xfrm). Still emit the item so its segments stay reviewable.
        return {"x": 0.0, "y": 0.0, "w": page_w, "h": page_h / 4}
    x, y, w, h = xf.apply(float(left), float(top), float(width), float(height))
    return {"x": _pt(x), "y": _pt(y), "w": _pt(w), "h": _pt(h)}


# ----- style inheritance ---------------------------------------------------


def _placeholder_chain(shape: Any, slide: Any) -> list[Any]:
    """txBody elements to consult for inherited paragraph styles, nearest first."""
    chain: list[Any] = []
    if not getattr(shape, "is_placeholder", False):
        return chain
    try:
        idx = shape.placeholder_format.idx
        layout_ph = slide.slide_layout.placeholders.get(idx=idx)
    except Exception:
        layout_ph = None
    if layout_ph is not None:
        tx = layout_ph._element.find(f"{{{P_NS}}}txBody")
        if tx is not None:
            chain.append(tx)
        try:
            ph_type = layout_ph._element.ph_type
            ph_type = _MASTER_PH_FOR.get(ph_type, ph_type)
            master_ph = slide.slide_layout.slide_master.placeholders.get(ph_type)
        except Exception:
            master_ph = None
        if master_ph is not None:
            tx = master_ph._element.find(f"{{{P_NS}}}txBody")
            if tx is not None:
                chain.append(tx)
    return chain


def _ph_type(shape: Any) -> str | None:
    ph = shape._element.find(f".//{{{P_NS}}}nvPr/{{{P_NS}}}ph")
    if ph is None:
        return None
    return str(ph.get("type", "body"))


def _master_style(shape: Any, slide: Any) -> Any | None:
    """The master's ``titleStyle`` / ``bodyStyle`` / ``otherStyle`` for this shape."""
    try:
        master = slide.slide_layout.slide_master._element
    except Exception:
        return None
    if not getattr(shape, "is_placeholder", False):
        name = "otherStyle"
    else:
        ph_type = _ph_type(shape)
        name = "titleStyle" if ph_type in _TITLE_PH_TYPES else "bodyStyle"
        if ph_type not in _TITLE_PH_TYPES and ph_type not in _BODY_PH_TYPES:
            name = "otherStyle"
    return master.find(f"{{{P_NS}}}txStyles/{{{P_NS}}}{name}")


class _StyleResolver:
    """Resolve size / bold / italic / align through the inheritance chain."""

    def __init__(self, shape: Any, slide: Any, tx_body: Any) -> None:
        self._own = tx_body
        self._chain = _placeholder_chain(shape, slide)
        self._master = _master_style(shape, slide)
        body_pr = tx_body.find(f"{{{A_NS}}}bodyPr") if tx_body is not None else None
        self.font_scale = 1.0
        self.anchor: str | None = None
        for tx in [tx_body, *self._chain]:
            if tx is None:
                continue
            bp = tx.find(f"{{{A_NS}}}bodyPr")
            if bp is None:
                continue
            if self.anchor is None and bp.get("anchor") in _ANCHOR_ATTR:
                self.anchor = _ANCHOR_ATTR[bp.get("anchor")]
        if body_pr is not None:
            fit = body_pr.find(f"{{{A_NS}}}normAutofit")
            if fit is not None and fit.get("fontScale"):
                self.font_scale = int(fit.get("fontScale")) / 100000.0

    def _lvl_pprs(self, para: Any) -> list[Any]:
        lvl = f"lvl{int(para.level) + 1}pPr"
        out: list[Any] = []
        ppr = para._p.find(f"{{{A_NS}}}pPr")
        if ppr is not None:
            out.append(ppr)
        for tx in [self._own, *self._chain]:
            if tx is None:
                continue
            el = tx.find(f"{{{A_NS}}}lstStyle/{{{A_NS}}}{lvl}")
            if el is not None:
                out.append(el)
        if self._master is not None:
            el = self._master.find(f"{{{A_NS}}}{lvl}")
            if el is not None:
                out.append(el)
        return out

    def _def_rpr_attr(self, para: Any, attr: str) -> str | None:
        for ppr in self._lvl_pprs(para):
            d = ppr.find(f"{{{A_NS}}}defRPr")
            if d is not None and d.get(attr) is not None:
                return str(d.get(attr))
        return None

    def align(self, para: Any) -> str:
        if para.alignment is not None:
            return _ALIGN.get(para.alignment, "left")
        for ppr in self._lvl_pprs(para):
            if ppr.get("algn") in _ALGN_ATTR:
                return _ALGN_ATTR[ppr.get("algn")]
        return "left"

    def size(self, run: Any, para: Any) -> float:
        if run.font.size is not None:
            pt = run.font.size.pt
        else:
            sz = self._def_rpr_attr(para, "sz")
            pt = int(sz) / 100.0 if sz is not None else DEFAULT_FONT_PT
        return round(float(pt) * self.font_scale, 2)

    def flag(self, value: bool | None, para: Any, attr: str) -> bool:
        if value is not None:
            return bool(value)
        inherited = self._def_rpr_attr(para, attr)
        return inherited in {"1", "true"}


def _rgb(color: Any) -> str | None:
    try:
        if color.type == MSO_COLOR_TYPE.RGB and color.rgb is not None:
            return f"#{color.rgb}"
    except Exception:
        return None
    return None


def _paragraphs(text_frame: Any, prefix: str, resolver: _StyleResolver) -> list[dict[str, Any]]:
    paragraphs: list[dict[str, Any]] = []
    for p_idx, para in enumerate(text_frame.paragraphs):
        runs = []
        for r_idx, run in enumerate(para.runs):
            runs.append(
                {
                    "segment_id": run_id(prefix, p_idx, r_idx),
                    "text": run.text,
                    "size": resolver.size(run, para),
                    "bold": resolver.flag(run.font.bold, para, "b"),
                    "italic": resolver.flag(run.font.italic, para, "i"),
                    "color": _rgb(run.font.color),
                    "font": run.font.name,
                }
            )
        paragraphs.append({"align": resolver.align(para), "runs": runs})
    return paragraphs


def _textbox(
    box: dict[str, float], text_frame: Any, prefix: str, shape: Any, slide: Any
) -> dict[str, Any]:
    resolver = _StyleResolver(shape, slide, text_frame._txBody)
    item: dict[str, Any] = {"kind": "textbox", **box}
    item["paragraphs"] = _paragraphs(text_frame, prefix, resolver)
    item["anchor"] = resolver.anchor or "top"
    return item


def _shape_fill(shape: Any) -> str | None:
    try:
        fill = shape.fill
        if fill.type == MSO_FILL.SOLID:
            return _rgb(fill.fore_color)
    except Exception:
        return None
    return None


def _cell_fill(cell: Any) -> str | None:
    try:
        if cell.fill.type == MSO_FILL.SOLID:
            return _rgb(cell.fill.fore_color)
    except Exception:
        return None
    return None


def _table_items(
    shape: Any, prefix: str, box: dict[str, float], slide: Any, xf: Transform
) -> list[dict[str, Any]]:
    table = shape.table
    col_w = [_pt(float(c.width) * xf.sx) for c in table.columns]
    row_h = [_pt(float(r.height) * xf.sy) for r in table.rows]
    items: list[dict[str, Any]] = []
    y = box["y"]
    for r_idx, row in enumerate(table.rows):
        x = box["x"]
        for c_idx, cell in enumerate(row.cells):
            w, h = col_w[c_idx], row_h[r_idx]
            if cell.is_merge_origin:
                w = sum(col_w[c_idx : c_idx + cell.span_width])
                h = sum(row_h[r_idx : r_idx + cell.span_height])
            cell_box = {"x": round(x, 2), "y": round(y, 2), "w": round(w, 2), "h": round(h, 2)}
            if not cell.is_spanned:
                fill = _cell_fill(cell) or (_TABLE_HEADER_FILL if r_idx == 0 else _TABLE_BODY_FILL)
                items.append({"kind": "shape", **cell_box, "fill": fill})
            cell_prefix = f"{prefix}/row{r_idx}/col{c_idx}"
            resolver = _StyleResolver(shape, slide, cell.text_frame._txBody)
            items.append(
                {
                    "kind": "textbox",
                    **cell_box,
                    "paragraphs": _paragraphs(cell.text_frame, cell_prefix, resolver),
                    "anchor": "top",
                }
            )
            x += col_w[c_idx]
        y += row_h[r_idx]
    return items


def _image_item(shape: Any, prefix: str, box: dict[str, float]) -> dict[str, Any]:
    src: str | None = None
    try:
        image = shape.image
        if image.content_type in _BROWSER_IMAGE_TYPES:
            b64 = base64.b64encode(image.blob).decode("ascii")
            src = f"data:{image.content_type};base64,{b64}"
    except Exception:
        src = None
    item: dict[str, Any] = {"kind": "image", **box, "src": src}
    text = alt_text(shape)
    item["alt"] = {"segment_id": f"{prefix}/alt", "text": text} if text.strip() else None
    return item


def _graphic_kind(shape: Any) -> str | None:
    gd = shape._element.find(f".//{{{A_NS}}}graphicData")
    if gd is None:
        return None
    return _GRAPHIC_DATA_KINDS.get(str(gd.get("uri", "")))


def _background_fill(slide: Any) -> str | None:
    """First explicit solid sRGB background on slide → layout → master."""
    parts = [slide, slide.slide_layout, slide.slide_layout.slide_master]
    for part in parts:
        bg = part._element.find(f"{{{P_NS}}}cSld/{{{P_NS}}}bg")
        if bg is None:
            continue
        clr = bg.find(f"{{{P_NS}}}bgPr/{{{A_NS}}}solidFill/{{{A_NS}}}srgbClr")
        if clr is not None and clr.get("val"):
            return f"#{clr.get('val')}"
        return None  # a background is defined but not a plain sRGB fill
    return None


def _slide_items(
    slide: Any, s_idx: int, page_w: float, page_h: float, unscanned: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    bg = _background_fill(slide)
    if bg is not None:
        items.append({"kind": "shape", "x": 0.0, "y": 0.0, "w": page_w, "h": page_h, "fill": bg})

    for shape, prefix, xf in iter_slide_shapes(slide, s_idx):
        if is_group(shape):
            continue
        box = _box(shape, xf, page_w, page_h)
        name = shape.name or "shape"
        if shape.has_table:
            items.extend(_table_items(shape, prefix, box, slide, xf))
            continue
        kind = _graphic_kind(shape)
        if kind is not None:
            items.append({"kind": "shape", **box, "fill": _UNRENDERED_FILL})
            unscanned.append(
                {
                    "where": f"Slide {s_idx + 1}",
                    "what": f'{kind} "{name}": its text is not scanned',
                    "page": s_idx,
                }
            )
            continue
        if shape._element.find(f"{{{P_NS}}}nvPicPr") is not None:
            items.append(_image_item(shape, prefix, box))
            continue
        fill = _shape_fill(shape)
        if fill is not None:
            items.append({"kind": "shape", **box, "fill": fill})
        if shape.has_text_frame:
            items.append(_textbox(box, shape.text_frame, prefix, shape, slide))
    return items


def _notes(slide: Any, s_idx: int) -> list[dict[str, Any]] | None:
    frame = notes_text_frame(slide)
    if frame is None:
        return None
    resolver = _StyleResolver(slide.notes_slide.notes_placeholder, slide.notes_slide, frame._txBody)
    # Notes are shown as a readable strip under the slide, not at their
    # printed size; normalise the size so a 12pt note isn't microscopic.
    paragraphs = _paragraphs(frame, f"slide{s_idx}/notes", resolver)
    for para in paragraphs:
        for run in para["runs"]:
            run["size"] = 12.0
    return paragraphs


def _master_text_unscanned(prs: Any) -> list[dict[str, Any]]:
    """Non-placeholder text on used layouts/masters is painted on slides but never scanned."""
    out: list[dict[str, Any]] = []
    seen: set[int] = set()
    for slide in prs.slides:
        layout = slide.slide_layout
        for part, label in ((layout, "Slide layout"), (layout.slide_master, "Slide master")):
            if id(part._element) in seen:
                continue
            seen.add(id(part._element))
            count = 0
            for shape in part.shapes:
                if getattr(shape, "is_placeholder", False):
                    continue
                if shape.has_text_frame and shape.text_frame.text.strip():
                    count += 1
            if count:
                out.append(
                    {
                        "where": f'{label} "{part.name}"',
                        "what": f"{count} text shape(s) shown on slides are not scanned",
                        "page": None,
                    }
                )
    return out


def _comments_unscanned(prs: Any) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for s_idx, slide in enumerate(prs.slides):
        if any(rel.reltype.endswith("/comments") for rel in slide.part.rels.values()):
            out.append(
                {
                    "where": f"Slide {s_idx + 1}",
                    "what": "Reviewer comments are not scanned; they are removed from the output",
                    "page": s_idx,
                }
            )
    return out


def _properties_unscanned(prs: Any) -> list[dict[str, Any]]:
    cp = prs.core_properties
    fields = [
        label
        for label, value in (
            ("author", cp.author),
            ("last modified by", cp.last_modified_by),
            ("title", cp.title),
            ("subject", cp.subject),
            ("keywords", cp.keywords),
            ("comments", cp.comments),
        )
        if value and str(value).strip()
    ]
    if not fields:
        return []
    return [
        {
            "where": "Document properties",
            "what": f"{', '.join(fields)} are not scanned; they are blanked in the output",
            "page": None,
        }
    ]


def build_layout(source: Path | BytesIO) -> dict[str, Any]:
    """Return the layout-contract payload for a .pptx file or byte stream."""
    prs = Presentation(source if isinstance(source, BytesIO) else str(source))
    page_w = _pt(float(prs.slide_width or 9144000))
    page_h = _pt(float(prs.slide_height or 6858000))
    unscanned: list[dict[str, Any]] = []
    pages: list[dict[str, Any]] = []
    for s_idx, slide in enumerate(prs.slides):
        items = _slide_items(slide, s_idx, page_w, page_h, unscanned)
        pages.append(
            {
                "index": s_idx,
                "width": page_w,
                "height": page_h,
                "items": items,
                "notes": _notes(slide, s_idx),
            }
        )
    unscanned.extend(_comments_unscanned(prs))
    unscanned.extend(_master_text_unscanned(prs))
    unscanned.extend(_properties_unscanned(prs))
    return {"format": "pptx", "pages": pages, "unscanned": unscanned}


def layout_segment_ids(layout: dict[str, Any]) -> list[str]:
    """Every segment id a layout renders, in render order (test/diagnostic helper)."""
    ids: list[str] = []
    for page in layout["pages"]:
        for item in page["items"]:
            if item["kind"] == "textbox":
                ids.extend(r["segment_id"] for p in item["paragraphs"] for r in p["runs"])
            elif item["kind"] == "image" and item.get("alt"):
                ids.append(item["alt"]["segment_id"])
        for para in page.get("notes") or []:
            ids.extend(r["segment_id"] for r in para["runs"])
    return ids
