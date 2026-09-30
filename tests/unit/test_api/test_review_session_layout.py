"""``GET /review-sessions/<id>/layout`` + a real-pptx commit (Phase 3.5 WS1).

Uses the real pptx Reader/Writer and the real Presidio anonymizer (the
``replace`` operator); only the analyzer is faked, with a name list, so
the test needs no spaCy model.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation
from sanctum.anonymizer.adapter import PresidioAnonymizer
from sanctum.api.app import create_app
from sanctum.api.schemas import ReviewSessionLayoutResponse
from sanctum.core.engine import SanctumEngine
from sanctum.core.models import DetectionResult
from sanctum.core.review.store import SessionStore
from sanctum.documents.pptx_layout import layout_segment_ids

LOOPBACK = {"Host": "127.0.0.1:8765"}
AUTH = {"Authorization": "Bearer t"}
PEOPLE = (
    "Margaret Holloway",
    "Daniel Okafor",
    "Jonas Albrecht",
    "Samuel Achterberg",
    "Olivia Brandt",
)


class _NameAnalyzer:
    def analyze(
        self,
        text: str,
        language: str = "en",
        entities: list[str] | None = None,
        score_threshold: float | None = None,
    ) -> list[DetectionResult]:
        out = []
        for name in PEOPLE:
            for m in re.finditer(re.escape(name), text):
                out.append(
                    DetectionResult(
                        entity_type="PERSON",
                        start=m.start(),
                        end=m.end(),
                        score=0.9,
                        text_span=name,
                    )
                )
        return sorted(out, key=lambda d: d.start)


@pytest.fixture()
def client(tmp_path: Path) -> Any:
    engine = SanctumEngine(analyzer=_NameAnalyzer(), anonymizer=PresidioAnonymizer())  # type: ignore[arg-type]
    store = SessionStore(root=tmp_path / "sessions")
    app = create_app(token="t", host="127.0.0.1", port=8765, engine=engine, session_store=store)
    return app.test_client()


def _create(client: Any, path: Path) -> dict[str, Any]:
    r = client.post(
        "/review-sessions",
        headers={**LOOPBACK, **AUTH},
        json={"input_path": str(path), "default_operator": "replace"},
    )
    assert r.status_code == 201, r.get_json()
    return r.get_json()


def test_layout_matches_session_segments(client: Any, rich_pptx: Path) -> None:
    session = _create(client, rich_pptx)
    r = client.get(f"/review-sessions/{session['id']}/layout", headers={**LOOPBACK, **AUTH})
    assert r.status_code == 200, r.get_json()
    body = r.get_json()
    ReviewSessionLayoutResponse.model_validate(body)
    assert body["format"] == "pptx"
    assert layout_segment_ids(body) == [s["id"] for s in session["segments"]]
    anchors = {p["segment_anchor"] for p in session["proposals"]}
    assert anchors <= set(layout_segment_ids(body))
    # notes and alt-text produced proposals too
    assert "slide0/notes/p0/r0" in anchors
    assert "slide2/shape1/alt" in anchors


def test_layout_errors(client: Any, rich_pptx: Path, tmp_path: Path) -> None:
    assert (
        client.get("/review-sessions/nope/layout", headers={**LOOPBACK, **AUTH}).status_code == 404
    )
    session = _create(client, rich_pptx)
    sid = session["id"]
    assert client.get(f"/review-sessions/{sid}/layout", headers=LOOPBACK).status_code == 401
    assert client.delete(f"/review-sessions/{sid}", headers={**LOOPBACK, **AUTH}).status_code == 204
    assert (
        client.get(f"/review-sessions/{sid}/layout", headers={**LOOPBACK, **AUTH}).status_code
        == 410
    )


def test_layout_unsupported_format_is_415(client: Any, tmp_path: Path) -> None:
    import docx

    d = docx.Document()
    d.add_paragraph("Margaret Holloway")
    path = tmp_path / "x.docx"
    d.save(str(path))
    session = _create(client, path)
    r = client.get(f"/review-sessions/{session['id']}/layout", headers={**LOOPBACK, **AUTH})
    assert r.status_code == 415


def test_commit_replaces_accepted_names_everywhere(
    client: Any, rich_pptx: Path, tmp_path: Path
) -> None:
    session = _create(client, rich_pptx)
    sid = session["id"]
    rejected = next(p for p in session["proposals"] if p["original"] == "Jonas Albrecht")
    for p in session["proposals"]:
        status = "reject" if p is rejected else "accept"
        r = client.patch(
            f"/review-sessions/{sid}/decisions/{p['detection_id']}",
            headers={**LOOPBACK, **AUTH},
            json={"status": status},
        )
        assert r.status_code == 200
    out = tmp_path / "out.pptx"
    r = client.post(
        f"/review-sessions/{sid}/commit",
        headers={**LOOPBACK, **AUTH},
        json={"output_path": str(out), "attested": True},
    )
    assert r.status_code == 200, r.get_json()

    prs = Presentation(str(out))
    texts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            for sub in [shape, *getattr(shape, "shapes", [])]:
                for leaf in [sub, *getattr(sub, "shapes", [])]:
                    if leaf.has_text_frame:
                        texts.append(leaf.text_frame.text)
            if shape.has_table:
                texts.extend(c.text for row in shape.table.rows for c in row.cells)
            c_nv_pr = shape._element.find(
                "{http://schemas.openxmlformats.org/presentationml/2006/main}nvPicPr/"
                "{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr"
            )
            if c_nv_pr is not None:
                texts.append(c_nv_pr.get("descr", ""))
        if slide.has_notes_slide:
            texts.append(slide.notes_slide.notes_text_frame.text)
    blob = "\n".join(texts)
    for name in PEOPLE:
        if name == "Jonas Albrecht":
            assert name in blob  # rejected → kept
        else:
            assert name not in blob, name
    assert "<PERSON>" in blob
    # formatting survived: the bold run is still bold and now holds the tag
    run = prs.slides[0].placeholders[1].text_frame.paragraphs[0].runs[1]
    assert run.font.bold is True and run.text == "<PERSON>"
