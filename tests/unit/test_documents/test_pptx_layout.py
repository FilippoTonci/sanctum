"""Unit tests for the pptx review layout (Phase 3.5 WS1.3)."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt
from sanctum.core.exceptions import UnsupportedDocumentFormatError
from sanctum.documents.layout import build_layout as dispatch_layout
from sanctum.documents.layout import supports_layout
from sanctum.documents.pptx_adapter import Reader
from sanctum.documents.pptx_layout import build_layout, layout_segment_ids

MEMO = Path("tests/fixtures/office/internal_memo.pptx")


def _textboxes(page: dict) -> list[dict]:
    return [i for i in page["items"] if i["kind"] == "textbox"]


@pytest.mark.parametrize("which", ["memo", "rich"])
def test_layout_segment_ids_equal_reader_segment_ids(which: str, rich_pptx: Path) -> None:
    path = MEMO if which == "memo" else rich_pptx
    layout = build_layout(path)
    doc = Reader().read(path)
    assert layout_segment_ids(layout) == [s.id for s in doc.segments]
    texts = {s.id: s.text for s in doc.segments}
    for page in layout["pages"]:
        for item in _textboxes(page):
            for para in item["paragraphs"]:
                for run in para["runs"]:
                    assert run["text"] == texts[run["segment_id"]]


def test_pages_are_in_points(rich_pptx: Path) -> None:
    layout = build_layout(rich_pptx)
    assert layout["format"] == "pptx"
    assert [p["index"] for p in layout["pages"]] == [0, 1, 2, 3]
    assert layout["pages"][0]["width"] == 720.0
    assert layout["pages"][0]["height"] == 540.0


def test_group_children_are_placed_on_the_slide(rich_pptx: Path) -> None:
    page = build_layout(rich_pptx)["pages"][2]
    card_fill = next(i for i in page["items"] if i["kind"] == "shape")
    assert card_fill["fill"] == "#123456"
    assert (card_fill["x"], card_fill["y"], card_fill["w"], card_fill["h"]) == (
        72.0,
        72.0,
        216.0,
        72.0,
    )
    deputy = next(
        i
        for i in _textboxes(page)
        if i["paragraphs"][0]["runs"][0]["text"] == "Deputy: Olivia Brandt"
    )
    assert (deputy["x"], deputy["y"]) == (72.0, 180.0)
    card_text = next(
        i for i in _textboxes(page) if i["paragraphs"][0]["runs"][0]["text"].startswith("GC:")
    )
    assert card_text["paragraphs"][0]["runs"][0]["size"] == 20.0


def test_scaled_group_transform(tmp_path: Path) -> None:
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    grp = s.shapes.add_group_shape()
    grp.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text_frame.text = "x"
    xfrm = grp._element.grpSpPr.xfrm
    # Squash the group to half size, keep child space: children scale by 0.5.
    xfrm.ext.cx = xfrm.ext.cx // 2
    xfrm.ext.cy = xfrm.ext.cy // 2
    path = tmp_path / "g.pptx"
    prs.save(str(path))
    box = _textboxes(build_layout(path)["pages"][0])[0]
    assert (box["x"], box["y"], box["w"], box["h"]) == (72.0, 72.0, 72.0, 36.0)


def test_table_cells_become_positioned_textboxes(rich_pptx: Path) -> None:
    page = build_layout(rich_pptx)["pages"][1]
    boxes = _textboxes(page)
    assert len(boxes) == 4
    jonas = next(b for b in boxes if b["paragraphs"][0]["runs"][0]["text"] == "Jonas Albrecht")
    assert jonas["x"] == 72.0
    assert jonas["y"] > 72.0
    assert sum(1 for i in page["items"] if i["kind"] == "shape") == 4


def test_picture_with_alt_text(rich_pptx: Path) -> None:
    page = build_layout(rich_pptx)["pages"][2]
    image = next(i for i in page["items"] if i["kind"] == "image")
    assert image["src"].startswith("data:image/png;base64,")
    assert image["alt"] == {"segment_id": "slide2/shape1/alt", "text": "Photo of Samuel Achterberg"}


def test_notes_and_inherited_title_styles(rich_pptx: Path) -> None:
    pages = build_layout(rich_pptx)["pages"]
    assert pages[0]["notes"][0]["runs"][0]["segment_id"] == "slide0/notes/p0/r0"
    assert pages[1]["notes"] is None
    title = _textboxes(pages[0])[0]
    assert title["paragraphs"][0]["align"] == "center"
    assert title["paragraphs"][0]["runs"][0]["size"] == 44.0  # from the master titleStyle
    sub_runs = _textboxes(pages[0])[1]["paragraphs"][0]["runs"]
    assert [r["bold"] for r in sub_runs] == [False, True]


def test_background_and_unscanned_chart(rich_pptx: Path) -> None:
    layout = build_layout(rich_pptx)
    first = layout["pages"][3]["items"][0]
    assert first == {"kind": "shape", "x": 0.0, "y": 0.0, "w": 720.0, "h": 540.0, "fill": "#FAF6EE"}
    charts = [u for u in layout["unscanned"] if "Chart" in u["what"]]
    assert charts and charts[0]["page"] == 3 and charts[0]["where"] == "Slide 4"


def test_autofit_font_scale_applies(tmp_path: Path) -> None:
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tf = s.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text_frame
    tf.text = "shrunk"
    tf.paragraphs[0].runs[0].font.size = Pt(40)
    body_pr = tf._txBody.bodyPr
    from lxml import etree

    fit = etree.SubElement(
        body_pr, "{http://schemas.openxmlformats.org/drawingml/2006/main}normAutofit"
    )
    fit.set("fontScale", "50000")
    path = tmp_path / "fit.pptx"
    prs.save(str(path))
    run = _textboxes(build_layout(path)["pages"][0])[0]["paragraphs"][0]["runs"][0]
    assert run["size"] == 20.0


def test_master_text_and_properties_are_reported(tmp_path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # MasterShapes has no add_textbox: build on the slide, move to the master.
    tb = slide.shapes.add_textbox(Inches(0), Inches(7), Inches(3), Inches(0.4))
    tb.text_frame.text = "Confidential - Holloway & Finch"
    prs.slide_masters[0].shapes._spTree.append(tb._element)
    prs.core_properties.author = "Someone"
    path = tmp_path / "m.pptx"
    prs.save(str(path))
    unscanned = build_layout(path)["unscanned"]
    assert any(u["where"].startswith("Slide master") for u in unscanned)
    assert any(u["where"] == "Document properties" and "author" in u["what"] for u in unscanned)


def test_dispatch_accepts_bytes_and_rejects_other_formats(rich_pptx: Path) -> None:
    assert supports_layout("pptx")
    assert not supports_layout("docx")
    layout = dispatch_layout("pptx", rich_pptx.read_bytes())
    assert len(layout["pages"]) == 4
    assert build_layout(BytesIO(rich_pptx.read_bytes()))["pages"] == layout["pages"]
    with pytest.raises(UnsupportedDocumentFormatError):
        dispatch_layout("docx", b"")
