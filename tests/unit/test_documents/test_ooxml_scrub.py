"""Word and PowerPoint output carries no hidden identifying data (E7).

Author, last-modified-by and the other text core properties, the app
properties' Company/Manager, every comment part and the package thumbnail
(a picture of the unredacted first page) are removed when the writer saves.
Both writers also implement ``extract_text`` so the engine's post-write
leak check covers .docx and .pptx. Every name below is invented.
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import docx
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as PPTX_RT
from pptx.opc.package import Part as PptxPart
from pptx.opc.packuri import PackURI
from pptx.util import Inches
from sanctum.core.leak_check import find_surviving_originals
from sanctum.core.protocols import OutputTextExtractor
from sanctum.documents import docx_adapter, pptx_adapter

NAME = "Jennifer Martin"
COMPANY = "Northgate Legal LLP"

_CORE_FIELDS = (
    "author",
    "last_modified_by",
    "title",
    "subject",
    "keywords",
    "comments",
    "category",
)


def _set_core(cp: object) -> None:
    for field in _CORE_FIELDS:
        setattr(cp, field, NAME)


def _core_values(cp: object) -> tuple[str, ...]:
    return tuple(getattr(cp, field) for field in _CORE_FIELDS)


def _add_app_identity(path: Path) -> None:
    """Rewrite docProps/app.xml so it names a Company and a Manager."""
    with zipfile.ZipFile(path) as z:
        entries = {n: z.read(n) for n in z.namelist()}
    app = entries["docProps/app.xml"].decode()
    extra = f"<Company>{COMPANY}</Company><Manager>{NAME}</Manager>"
    app = re.sub(r"</Properties>\s*$", extra + "</Properties>", app)
    entries["docProps/app.xml"] = app.encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in entries.items():
            z.writestr(name, data)


def _all_bytes(path: Path) -> bytes:
    with zipfile.ZipFile(path) as z:
        return b"".join(z.read(n) for n in z.namelist())


def _names(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as z:
        return z.namelist()


# --------------------------------------------------------------------- docx


def make_docx_with_hidden_data(path: Path) -> Path:
    d = docx.Document()
    _set_core(d.core_properties)
    para = d.add_paragraph()
    para.add_run("The client ")
    run = para.add_run("signed")
    para.add_run(" the agreement.")
    d.add_comment(run, text=f"Ask {NAME} to countersign", author=NAME, initials="JM")
    d.save(str(path))
    _add_app_identity(path)
    return path


def roundtrip_docx(src: Path, out: Path) -> Path:
    doc = docx_adapter.Reader().read(src)
    docx_adapter.Writer().write(doc, out)
    return out


def test_docx_fixture_really_carries_hidden_data(tmp_path: Path) -> None:
    src = make_docx_with_hidden_data(tmp_path / "in.docx")
    assert any(n.startswith("word/comments") for n in _names(src))
    assert NAME.encode() in _all_bytes(src)
    assert COMPANY.encode() in _all_bytes(src)
    assert "docProps/thumbnail.jpeg" in _names(src)


def test_docx_output_has_no_identifying_properties_or_comments(tmp_path: Path) -> None:
    src = make_docx_with_hidden_data(tmp_path / "in.docx")
    out = roundtrip_docx(src, tmp_path / "out.docx")
    d = docx.Document(str(out))
    assert _core_values(d.core_properties) == ("",) * len(_CORE_FIELDS)
    names = _names(out)
    assert not any(n.startswith("word/comments") for n in names)
    assert not any(n.startswith("docProps/thumbnail") for n in names)
    blob = _all_bytes(out)
    assert NAME.encode() not in blob
    assert COMPANY.encode() not in blob
    for tag in (b"commentRangeStart", b"commentRangeEnd", b"commentReference"):
        assert tag not in blob


def test_docx_output_keeps_the_body_text(tmp_path: Path) -> None:
    src = make_docx_with_hidden_data(tmp_path / "in.docx")
    out = roundtrip_docx(src, tmp_path / "out.docx")
    assert docx.Document(str(out)).paragraphs[0].text == "The client signed the agreement."


def test_docx_writer_is_an_output_text_extractor() -> None:
    assert isinstance(docx_adapter.Writer(), OutputTextExtractor)


def test_docx_extract_text_joins_runs_of_one_paragraph(tmp_path: Path) -> None:
    d = docx.Document()
    para = d.add_paragraph()
    for i, text in enumerate(["Dear Jen", "nifer Mar", "tin, thanks."]):
        para.add_run(text).bold = i % 2 == 1
    d.add_paragraph("Second paragraph.")
    path = tmp_path / "split.docx"
    d.save(str(path))
    text = docx_adapter.Writer().extract_text(path)
    assert text.splitlines() == ["Dear Jennifer Martin, thanks.", "Second paragraph."]
    assert find_surviving_originals(text, [NAME]) == [NAME]


def test_docx_extract_text_includes_tables_headers_and_footers(tmp_path: Path) -> None:
    d = docx.Document()
    d.add_paragraph("Body.")
    d.add_table(rows=1, cols=1).cell(0, 0).text = "In a table"
    section = d.sections[0]
    section.header.paragraphs[0].text = "Header line"
    section.footer.paragraphs[0].text = "Footer line"
    section.header.add_table(rows=1, cols=1, width=Inches(2)).cell(0, 0).text = "Header cell"
    path = tmp_path / "hf.docx"
    d.save(str(path))
    lines = docx_adapter.Writer().extract_text(path).splitlines()
    for expected in ("Body.", "In a table", "Header line", "Footer line", "Header cell"):
        assert expected in lines


# --------------------------------------------------------------------- pptx

_MODERN_COMMENTS_RT = "http://schemas.microsoft.com/office/2018/10/relationships/comments"
_MODERN_AUTHORS_RT = "http://schemas.microsoft.com/office/2018/10/relationships/authors"
_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_P188 = "http://schemas.microsoft.com/office/powerpoint/2018/8/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _xml_part(package: object, partname: str, content_type: str, body: str) -> PptxPart:
    blob = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>' + body).encode()
    return PptxPart(PackURI(partname), content_type, package, blob)  # type: ignore[arg-type]


def make_pptx_with_hidden_data(path: Path) -> Path:
    prs = Presentation()
    _set_core(prs.core_properties)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.text_frame.text = "Quarterly figures"
    package = prs.part.package

    # Legacy comments: ppt/commentAuthors.xml + ppt/comments/comment1.xml.
    authors = _xml_part(
        package,
        "/ppt/commentAuthors.xml",
        "application/vnd.openxmlformats-officedocument.presentationml.commentAuthors+xml",
        f'<p:cmAuthorLst xmlns:p="{_P}"><p:cmAuthor id="0" name="{NAME}" initials="JM"'
        ' lastIdx="1" clrIdx="0"/></p:cmAuthorLst>',
    )
    prs.part.relate_to(authors, PPTX_RT.COMMENT_AUTHORS)
    comment = _xml_part(
        package,
        "/ppt/comments/comment1.xml",
        "application/vnd.openxmlformats-officedocument.presentationml.comments+xml",
        f'<p:cmLst xmlns:p="{_P}"><p:cm authorId="0" idx="1"><p:pos x="10" y="10"/>'
        f"<p:text>Check with {NAME}</p:text></p:cm></p:cmLst>",
    )
    slide.part.relate_to(comment, PPTX_RT.COMMENTS)

    # Modern (threaded) comments: ppt/authors.xml + ppt/comments/modernComment_*.xml,
    # referenced from the slide's extLst by r:id.
    modern_authors = _xml_part(
        package,
        "/ppt/authors.xml",
        "application/vnd.ms-powerpoint.authors+xml",
        f'<p188:authorLst xmlns:p188="{_P188}"><p188:author id="{{00000000-0000-0000-0000-'
        f'000000000001}}" name="{NAME}" initials="JM" userId="{NAME}" providerId="None"/>'
        "</p188:authorLst>",
    )
    prs.part.relate_to(modern_authors, _MODERN_AUTHORS_RT)
    modern = _xml_part(
        package,
        "/ppt/comments/modernComment_100_1.xml",
        "application/vnd.ms-powerpoint.comments+xml",
        f'<p188:cmLst xmlns:p188="{_P188}" xmlns:a="http://schemas.openxmlformats.org/'
        'drawingml/2006/main"><p188:cm id="{00000000-0000-0000-0000-000000000002}" '
        'authorId="{00000000-0000-0000-0000-000000000001}" created="2026-01-01T00:00:00Z">'
        f"<p188:txBody><a:bodyPr/><a:p><a:r><a:t>Ask {NAME}</a:t></a:r></a:p></p188:txBody>"
        "</p188:cm></p188:cmLst>",
    )
    r_id = slide.part.relate_to(modern, _MODERN_COMMENTS_RT)
    from lxml import etree

    ext_lst = etree.SubElement(slide._element, f"{{{_P}}}extLst")
    ext = etree.SubElement(ext_lst, f"{{{_P}}}ext", uri="{6950BFC3-D8DA-4A85-94F7-54DA5524770B}")
    etree.SubElement(ext, f"{{{_P188}}}commentRel", {f"{{{_R}}}id": r_id})

    prs.save(str(path))
    _add_app_identity(path)
    return path


def roundtrip_pptx(src: Path, out: Path) -> Path:
    doc = pptx_adapter.Reader().read(src)
    pptx_adapter.Writer().write(doc, out)
    return out


def test_pptx_fixture_really_carries_hidden_data(tmp_path: Path) -> None:
    src = make_pptx_with_hidden_data(tmp_path / "in.pptx")
    names = _names(src)
    assert "ppt/commentAuthors.xml" in names
    assert "ppt/authors.xml" in names
    assert any(n.startswith("ppt/comments/") for n in names)
    assert "docProps/thumbnail.jpeg" in names
    assert NAME.encode() in _all_bytes(src)


def test_pptx_output_has_no_identifying_properties_or_comments(tmp_path: Path) -> None:
    src = make_pptx_with_hidden_data(tmp_path / "in.pptx")
    out = roundtrip_pptx(src, tmp_path / "out.pptx")
    prs = Presentation(str(out))
    assert _core_values(prs.core_properties) == ("",) * len(_CORE_FIELDS)
    names = _names(out)
    assert not any(n.startswith("ppt/comments/") for n in names)
    assert "ppt/commentAuthors.xml" not in names
    assert "ppt/authors.xml" not in names
    assert not any(n.startswith("docProps/thumbnail") for n in names)
    blob = _all_bytes(out)
    assert NAME.encode() not in blob
    assert COMPANY.encode() not in blob
    assert b"commentRel" not in blob
    assert prs.slides[0].shapes[0].text_frame.text == "Quarterly figures"


def test_pptx_writer_is_an_output_text_extractor() -> None:
    assert isinstance(pptx_adapter.Writer(), OutputTextExtractor)


def test_pptx_extract_text_joins_runs_and_reads_notes(tmp_path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    para = box.text_frame.paragraphs[0]
    for i, text in enumerate(["Dear Jen", "nifer Mar", "tin, thanks."]):
        run = para.add_run()
        run.text = text
        run.font.bold = i % 2 == 1
    slide.notes_slide.notes_text_frame.text = "Speaker note"
    path = tmp_path / "split.pptx"
    prs.save(str(path))
    text = pptx_adapter.Writer().extract_text(path)
    assert text.splitlines() == ["Dear Jennifer Martin, thanks.", "Speaker note"]
    assert find_surviving_originals(text, [NAME]) == [NAME]


# ----------------------------------------------------- package integrity (Ruling 16)
#
# Office is not available to open the scrubbed files, and neither is LibreOffice
# here, so check what Word / PowerPoint would trip over: the file re-opens, every
# internal relationship points at a part that exists, every r:id / r:embed a part
# uses is one of its relationships, and [Content_Types].xml declares no missing part.
# Opening one scrubbed .docx and .pptx in Office stays a manual pre-release check.

_PKG_REL = "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship"
_CT_OVERRIDE = "{http://schemas.openxmlformats.org/package/2006/content-types}Override"


def _assert_package_is_consistent(path: Path) -> None:
    from posixpath import dirname, join, normpath

    from lxml import etree

    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        parts = {n: z.read(n) for n in names}

    types = etree.fromstring(parts["[Content_Types].xml"])
    for override in types.iter(_CT_OVERRIDE):
        assert override.get("PartName").lstrip("/") in names, override.get("PartName")

    for rels_name in (n for n in names if n.endswith(".rels")):
        # word/_rels/document.xml.rels describes word/document.xml; _rels/.rels the package.
        source = rels_name.replace("_rels/", "").removesuffix(".rels")
        base = dirname(source)
        rel_ids = set()
        for rel in etree.fromstring(parts[rels_name]).iter(_PKG_REL):
            rel_ids.add(rel.get("Id"))
            if rel.get("TargetMode") == "External":
                continue
            target = rel.get("Target")
            resolved = (
                target.lstrip("/") if target.startswith("/") else normpath(join(base, target))
            )
            assert resolved in names, f"{rels_name}: {rel.get('Id')} -> {target} is missing"
        if source in parts and source.endswith(".xml"):
            used = set(
                re.findall(rb'\br:(?:id|embed|link|pict|dm|lo|qs|cs)="([^"]+)"', parts[source])
            )
            missing = {u.decode() for u in used} - rel_ids
            assert not missing, f"{source} refers to unknown relationships {missing}"


_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _add_link_and_tracked_change(path: Path) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import parse_xml

    d = docx.Document(str(path))
    rid = d.part.relate_to(f"mailto:{NAME}@example.com", RT.HYPERLINK, is_external=True)
    ns = f'xmlns:w="{_W_NS}" xmlns:r="{_R}"'
    p = (
        f'<w:p {ns}><w:hyperlink r:id="{rid}"><w:r><w:t>{NAME}</w:t></w:r></w:hyperlink>'
        f'<w:ins w:id="21" w:author="{NAME}" w:date="2026-01-01T00:00:00Z"><w:r>'
        f'<w:t xml:space="preserve"> added</w:t></w:r></w:ins>'
        f'<w:del w:id="22" w:author="{NAME}" w:date="2026-01-01T00:00:00Z"><w:r>'
        f"<w:delText>gone</w:delText></w:r></w:del></w:p>"
    )
    d.element.body[-1].addprevious(parse_xml(p))
    d.save(str(path))


def test_scrubbed_docx_reopens_with_no_dangling_parts(tmp_path: Path) -> None:
    src = make_docx_with_hidden_data(tmp_path / "in.docx")
    _add_link_and_tracked_change(src)
    _assert_package_is_consistent(src)  # the check itself accepts a sound input
    doc = docx_adapter.Reader().read(src)
    redacted = [
        s.model_copy(update={"text": "<PERSON>"}) if s.text == NAME else s for s in doc.segments
    ]
    assert redacted != doc.segments  # the link text really was redacted
    mutated = doc.model_copy(update={"segments": redacted})
    mutated.raw_handle = doc.raw_handle
    out = tmp_path / "out.docx"
    docx_adapter.Writer().write(mutated, out)
    _assert_package_is_consistent(out)
    reopened = docx.Document(str(out))
    assert "<PERSON> added" in [p.text for p in reopened.paragraphs]
    assert NAME.encode() not in _all_bytes(out)


def test_scrubbed_pptx_reopens_with_no_dangling_parts(tmp_path: Path) -> None:
    src = make_pptx_with_hidden_data(tmp_path / "in.pptx")
    _assert_package_is_consistent(src)
    out = roundtrip_pptx(src, tmp_path / "out.pptx")
    _assert_package_is_consistent(out)
    reopened = Presentation(str(out))
    assert reopened.slides[0].shapes[0].text_frame.text == "Quarterly figures"
