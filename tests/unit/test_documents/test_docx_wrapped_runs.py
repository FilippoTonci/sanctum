"""Runs inside w:hyperlink, w:ins, w:sdt and w:smartTag are read, redacted and leak-checked.

``Paragraph.runs`` only returns direct ``w:r``
children, so a name inside a mailto link, a tracked insertion or a content
control used to be neither detected nor leak-checked. Tracked changes are
accepted in the output, which also removes deleted text and revision
authors. Every name below is invented.
"""

from __future__ import annotations

import re
import struct
import zipfile
import zlib
from pathlib import Path

import docx
import pytest
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml
from sanctum.anonymizer.adapter import PresidioAnonymizer
from sanctum.core.engine import SanctumEngine
from sanctum.core.exceptions import LeakCheckError
from sanctum.core.leak_check import find_surviving_originals
from sanctum.core.models import DetectionResult, OperatorPolicy
from sanctum.documents.docx_adapter import Reader, Writer

NAME = "Jane Doe"
EMAIL = "jane.doe@example.com"
AUTHOR = "Rob Reviewer"
_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_NS = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
)
_DATE = 'w:date="2026-01-01T00:00:00Z"'
REPLACE = {"DEFAULT": OperatorPolicy(operator_name="replace")}


def _run(text: str) -> str:
    return f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r>'


def _fragments(rid: str) -> dict[str, str]:
    """Body-level XML for each wrapper kind; every one hides NAME or EMAIL."""
    return {
        "hyperlink": (
            f"<w:p {_NS}>{_run('Email ')}"
            f'<w:hyperlink r:id="{rid}">{_run(EMAIL)}</w:hyperlink>{_run(" today.")}</w:p>'
        ),
        "ins": (
            f"<w:p {_NS}>{_run('Signed by ')}"
            f'<w:ins w:id="11" w:author="{AUTHOR}" {_DATE}>{_run(NAME)}</w:ins></w:p>'
        ),
        "sdt_inline": (
            f"<w:p {_NS}>{_run('Client: ')}"
            f'<w:sdt><w:sdtPr><w:alias w:val="Client"/></w:sdtPr>'
            f"<w:sdtContent>{_run(NAME)}</w:sdtContent></w:sdt></w:p>"
        ),
        "smart_tag": (
            f"<w:p {_NS}>{_run('Contact ')}"
            f'<w:smartTag w:uri="urn:x" w:element="PersonName">{_run(NAME)}</w:smartTag></w:p>'
        ),
        "sdt_block": (
            f"<w:sdt {_NS}><w:sdtPr/><w:sdtContent>"
            f"<w:p>{_run('Party: ' + NAME)}</w:p></w:sdtContent></w:sdt>"
        ),
    }


_DELETION = (
    f"<w:p {_NS}>{_run('Witness: ')}"
    f'<w:del w:id="12" w:author="{AUTHOR}" {_DATE}>'
    f'<w:r><w:delText xml:space="preserve">{NAME}</w:delText></w:r></w:del>'
    f"{_run('none')}</w:p>"
)
_FORMAT_CHANGE = (
    f'<w:p {_NS}><w:r><w:rPr><w:b/><w:rPrChange w:id="13" w:author="{AUTHOR}" {_DATE}>'
    f"<w:rPr/></w:rPrChange></w:rPr><w:t>Bold now.</w:t></w:r></w:p>"
)


def make_docx(path: Path, kinds: list[str], *, deletion: bool = False) -> Path:
    d = docx.Document()
    d.add_paragraph(f"Dear {NAME}, see {EMAIL}.")
    rid = (
        d.part.relate_to(f"mailto:{EMAIL}", RT.HYPERLINK, is_external=True)
        if "hyperlink" in kinds
        else ""
    )
    body = d.element.body
    sect = body[-1]  # w:sectPr stays last
    fragments = _fragments(rid)
    xml = [fragments[k] for k in kinds]
    if deletion:
        xml += [_DELETION, _FORMAT_CHANGE]
    for fragment in xml:
        sect.addprevious(parse_xml(fragment))
    d.save(str(path))
    return path


class _FindAll:
    """Finds NAME and EMAIL wherever they appear (only text starting "Dear" if body_only)."""

    def __init__(self, body_only: bool = False) -> None:
        self.body_only = body_only

    def analyze(
        self,
        text: str,
        language: str = "en",
        entities: list[str] | None = None,
        score_threshold: float | None = None,
    ) -> list[DetectionResult]:
        if self.body_only and not text.startswith("Dear"):
            return []
        # Names found once are propagated to every block (core/propagation.py),
        # which would hide the miss body_only simulates; tag NAME with a type
        # propagation leaves alone so the leak check is what catches it.
        name_kind = "ID_NUMBER" if self.body_only else "PERSON"
        out = []
        for value, kind in ((NAME, name_kind), (EMAIL, "EMAIL_ADDRESS")):
            for m in re.finditer(re.escape(value), text):
                out.append(
                    DetectionResult(
                        entity_type=kind, start=m.start(), end=m.end(), score=0.9, text_span=value
                    )
                )
        return out


def _engine(body_only: bool = False) -> SanctumEngine:
    return SanctumEngine(analyzer=_FindAll(body_only), anonymizer=PresidioAnonymizer())


def _zip_bytes(path: Path) -> bytes:
    with zipfile.ZipFile(path) as z:
        return b"".join(z.read(n) for n in z.namelist())


ALL_KINDS = ["hyperlink", "ins", "sdt_inline", "smart_tag", "sdt_block"]


def test_reader_reads_runs_inside_wrappers(tmp_path: Path) -> None:
    src = make_docx(tmp_path / "in.docx", ALL_KINDS, deletion=True)
    segs = {s.id: s for s in Reader().read(src).segments}
    texts = [s.text for s in segs.values()]
    assert texts.count(NAME) == 3  # ins, inline sdt, smart tag
    assert EMAIL in texts  # hyperlink
    assert f"Party: {NAME}" in texts  # block-level content control
    # Plain runs keep the ids paragraph.runs would give them; wrapped runs get /n{k}.
    assert segs["body/p1/r0"].text == "Email "
    assert segs["body/p1/n0"].text == EMAIL
    assert segs["body/p1/r1"].text == " today."
    assert segs["body/p1/n0"].block == segs["body/p1/r0"].block == "body/p1"
    assert segs["body/sdt0/p0/r0"].text == f"Party: {NAME}"
    # Deleted text is not a segment: the output accepts the deletion.
    assert segs["body/p5/r0"].text == "Witness: "
    assert segs["body/p5/r1"].text == "none"
    assert not any(t == NAME for i, t in ((i, s.text) for i, s in segs.items()) if "p5" in i)


def test_plain_run_ids_are_unchanged(tmp_path: Path) -> None:
    d = docx.Document()
    para = d.add_paragraph()
    for i, text in enumerate(["one ", "two ", "three"]):
        para.add_run(text).bold = i % 2 == 1
    d.add_table(rows=1, cols=1).cell(0, 0).text = "cell"
    d.sections[0].header.paragraphs[0].text = "head"
    path = tmp_path / "plain.docx"
    d.save(str(path))
    ids = [s.id for s in Reader().read(path).segments]
    assert ids == [
        "body/p0/r0",
        "body/p0/r1",
        "body/p0/r2",
        "table/t0/row0/cell0/p0/r0",
        "hf/s0/header/p0/r0",
    ]


@pytest.mark.parametrize("kind", ALL_KINDS)
def test_wrapped_value_is_redacted_when_detected(tmp_path: Path, kind: str) -> None:
    src = make_docx(tmp_path / "in.docx", [kind])
    out = tmp_path / "out.docx"
    _engine().process_document(Reader(), Writer(), src, out, operator_policies=REPLACE)
    raw = _zip_bytes(out)
    assert NAME.encode() not in raw
    assert EMAIL.encode() not in raw  # also the mailto: relationship target
    docx.Document(str(out))  # still opens


@pytest.mark.parametrize("kind", ALL_KINDS)
def test_wrapped_value_missed_by_detection_fails_closed(tmp_path: Path, kind: str) -> None:
    src = make_docx(tmp_path / "in.docx", [kind])
    out = tmp_path / "out.docx"
    with pytest.raises(LeakCheckError) as exc_info:
        _engine(body_only=True).process_document(
            Reader(), Writer(), src, out, operator_policies=REPLACE
        )
    assert set(exc_info.value.leaks) & {NAME, EMAIL}
    assert not out.exists()


def test_tracked_deletion_and_authors_are_gone_from_the_output(tmp_path: Path) -> None:
    src = make_docx(tmp_path / "in.docx", ["ins"], deletion=True)
    out = tmp_path / "out.docx"
    _engine().process_document(Reader(), Writer(), src, out, operator_policies=REPLACE)
    raw = _zip_bytes(out)
    assert NAME.encode() not in raw
    assert AUTHOR.encode() not in raw
    for tag in (b"<w:del ", b"<w:ins ", b"w:delText", b"w:rPrChange"):
        assert tag not in raw
    texts = [p.text for p in docx.Document(str(out)).paragraphs]
    assert "Witness: none" in texts
    assert "Signed by <PERSON>" in texts
    assert "Bold now." in texts


def test_extract_text_reads_every_text_node_independently_of_the_reader(tmp_path: Path) -> None:
    src = make_docx(tmp_path / "in.docx", ALL_KINDS, deletion=True)
    lines = Writer().extract_text(src).splitlines()
    assert f"Email {EMAIL} today." in lines
    assert f"Signed by {NAME}" in lines
    assert f"Client: {NAME}" in lines
    assert f"Contact {NAME}" in lines
    assert f"Party: {NAME}" in lines
    assert f"Witness: {NAME}none" in lines  # deleted text still counts if it survives


def test_unchanged_runs_keep_inline_pictures(tmp_path: Path) -> None:
    # Setting Run.text clears the run's content; the writer now skips runs whose
    # text did not change, so a picture next to a redacted name survives.
    def chunk(kind: bytes, data: bytes) -> bytes:
        body = kind + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

    png = (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(b"\x00\xff\xff\xff"))
        + chunk(b"IEND", b"")
    )
    (tmp_path / "dot.png").write_bytes(png)
    d = docx.Document()
    d.add_paragraph(f"Dear {NAME},")
    d.add_paragraph().add_run().add_picture(str(tmp_path / "dot.png"))
    src = tmp_path / "in.docx"
    d.save(str(src))
    out = tmp_path / "out.docx"
    _engine().process_document(Reader(), Writer(), src, out, operator_policies=REPLACE)
    assert len(docx.Document(str(out)).inline_shapes) == 1


# ------------------------------------------- field codes and alt text

_FIELD = (
    f"<w:p {_NS}>{_run('Mail ')}"
    '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
    f'<w:r><w:instrText xml:space="preserve"> HYPERLINK "mailto:{EMAIL}" </w:instrText></w:r>'
    '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
    "{result}"
    '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
)
_FLD_SIMPLE = (
    f"<w:p {_NS}>{_run('Or ')}"
    f'<w:fldSimple w:instr=" HYPERLINK &quot;mailto:{EMAIL}&quot; ">{{result}}</w:fldSimple></w:p>'
)
_ALT = f"Passport photo of {NAME}"


def _png(path: Path) -> Path:
    def chunk(kind: bytes, data: bytes) -> bytes:
        body = kind + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

    path.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(b"\x00\xff\xff\xff"))
        + chunk(b"IEND", b"")
    )
    return path


def make_hidden_text_docx(path: Path, *, field_result: str, alt: bool = True) -> Path:
    """Body names NAME and EMAIL; a complex and a simple HYPERLINK field to EMAIL show
    ``field_result``; a picture's alt text names NAME."""
    d = docx.Document()
    d.add_paragraph(f"Dear {NAME}, see {EMAIL}.")
    sect = d.element.body[-1]
    sect.addprevious(parse_xml(_FIELD.replace("{result}", _run(field_result))))
    sect.addprevious(parse_xml(_FLD_SIMPLE.replace("{result}", _run(field_result))))
    if alt:
        shape = d.add_paragraph().add_run().add_picture(str(_png(path.with_suffix(".png"))))
        inline = shape._inline
        inline.docPr.set("descr", _ALT)
        inline.docPr.set("title", NAME)
        inline.docPr.set("name", f"{NAME}.png")
        cnvpr = inline.graphic.graphicData.pic.nvPicPr.cNvPr
        cnvpr.set("descr", _ALT)
        cnvpr.set("name", f"{NAME}.png")
    d.save(str(path))
    return path


def test_redacted_field_result_unwraps_the_field(tmp_path: Path) -> None:
    src = make_hidden_text_docx(tmp_path / "in.docx", field_result=EMAIL)
    out = tmp_path / "out.docx"
    _engine().process_document(Reader(), Writer(), src, out, operator_policies=REPLACE)
    raw = _zip_bytes(out)
    assert EMAIL.encode() not in raw
    for tag in (b"w:instrText", b"w:fldChar", b"w:fldSimple"):
        assert tag not in raw
    texts = [p.text for p in docx.Document(str(out)).paragraphs]
    assert "Mail <EMAIL_ADDRESS>" in texts
    assert "Or <EMAIL_ADDRESS>" in texts


def test_alt_text_naming_a_redacted_value_is_blanked(tmp_path: Path) -> None:
    src = make_hidden_text_docx(tmp_path / "in.docx", field_result=EMAIL)
    out = tmp_path / "out.docx"
    _engine().process_document(Reader(), Writer(), src, out, operator_policies=REPLACE)
    assert NAME.encode() not in _zip_bytes(out)
    written = docx.Document(str(out))
    assert len(written.inline_shapes) == 1
    doc_pr = written.inline_shapes[0]._inline.docPr
    assert doc_pr.get("name") == ""  # required by the schema: kept, emptied
    assert doc_pr.get("descr") == ""


def test_alt_text_is_blanked_on_the_review_commit_path_too(tmp_path: Path) -> None:
    from sanctum.core.models import ProposalDecision
    from sanctum.core.review.session import add_decision
    from sanctum.core.review.store import SessionStore

    store = SessionStore(root=tmp_path / "sessions")
    src = make_hidden_text_docx(tmp_path / "in.docx", field_result=EMAIL)
    engine = _engine()
    session = engine.create_review_session(
        Reader(), src, default_operator="replace", session_store=store
    )
    with store.locked(session.id):
        s = store.load(session.id)
        for p in s.proposals:
            add_decision(s, ProposalDecision(proposal_id=p.detection_id, status="accept"))
        store.save(s)
    out = tmp_path / "out.docx"
    engine.commit_review_session(Reader(), Writer(), session.id, out, store)
    raw = _zip_bytes(out)
    assert NAME.encode() not in raw
    assert EMAIL.encode() not in raw


def test_field_codes_and_alt_text_are_leak_checked(tmp_path: Path) -> None:
    # If the writer were bypassed, extract_text alone must surface both: here
    # neither value appears in any w:t outside the body line we drop.
    src = make_hidden_text_docx(tmp_path / "in.docx", field_result="click here")
    text = Writer().extract_text(src)
    hidden = "\n".join(line for line in text.splitlines() if not line.startswith("Dear "))
    assert set(find_surviving_originals(hidden, [NAME, EMAIL])) == {NAME, EMAIL}


def test_field_target_behind_unredacted_text_fails_closed(tmp_path: Path) -> None:
    # The field's visible text ("click here") is not redacted, so its code keeps
    # the email; the leak check must refuse the output.
    src = make_hidden_text_docx(tmp_path / "in.docx", field_result="click here", alt=False)
    out = tmp_path / "out.docx"
    with pytest.raises(LeakCheckError) as exc_info:
        _engine().process_document(Reader(), Writer(), src, out, operator_policies=REPLACE)
    assert EMAIL in exc_info.value.leaks
    assert not out.exists()


def _shared_rel_docx(path: Path, second_text: str) -> Path:
    d = docx.Document()
    rid = d.part.relate_to(f"mailto:{EMAIL}", RT.HYPERLINK, is_external=True)
    sect = d.element.body[-1]
    for text in (EMAIL, second_text):
        sect.addprevious(
            parse_xml(f'<w:p {_NS}><w:hyperlink r:id="{rid}">{_run(text)}</w:hyperlink></w:p>')
        )
    d.save(str(path))
    return path


def test_two_redacted_links_sharing_a_relationship(tmp_path: Path) -> None:
    src = _shared_rel_docx(tmp_path / "in.docx", EMAIL)
    out = tmp_path / "out.docx"
    _engine().process_document(Reader(), Writer(), src, out, operator_policies=REPLACE)
    assert EMAIL.encode() not in _zip_bytes(out)


def test_unlinking_one_link_keeps_a_relationship_another_link_uses(tmp_path: Path) -> None:
    src = _shared_rel_docx(tmp_path / "in.docx", "write to us")
    doc = Reader().read(src)
    mutated = doc.model_copy(
        update={
            "segments": [
                s.model_copy(update={"text": "<EMAIL_ADDRESS>"}) if s.text == EMAIL else s
                for s in doc.segments
            ]
        }
    )
    mutated.raw_handle = doc.raw_handle
    out = tmp_path / "out.docx"
    Writer().write(mutated, out)
    written = docx.Document(str(out))
    used = written.element.body.xpath("//w:hyperlink/@r:id")
    assert len(used) == 1
    assert used[0] in written.part.rels  # no dangling reference


# --------------------------------- VML alt text and internal-link tooltips

_VML_NS = (
    'xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office"'
)
_VML = (
    f"<w:p {_NS} {_VML_NS}><w:r><w:pict>"
    f'<v:shape id="pic1" style="width:10pt;height:10pt" alt="Photo of {NAME}">'
    f'<v:imagedata o:title="{NAME}"/></v:shape></w:pict></w:r></w:p>'
)
_ANCHOR_LINK = (
    f'<w:p {_NS}><w:hyperlink w:anchor="terms" w:tooltip="Ask {NAME}">'
    f"{_run('see the terms')}</w:hyperlink></w:p>"
)


def _docx_with(path: Path, *fragments: str) -> Path:
    d = docx.Document()
    d.add_paragraph(f"Dear {NAME},")
    for fragment in fragments:
        d.element.body[-1].addprevious(parse_xml(fragment))
    d.save(str(path))
    return path


def test_vml_alt_text_naming_a_redacted_value_is_blanked(tmp_path: Path) -> None:
    src = _docx_with(tmp_path / "in.docx", _VML)
    out = tmp_path / "out.docx"
    _engine().process_document(Reader(), Writer(), src, out, operator_policies=REPLACE)
    assert NAME.encode() not in _zip_bytes(out)
    body = docx.Document(str(out)).element.body
    [shape] = body.iter("{urn:schemas-microsoft-com:vml}shape")
    assert shape.get("alt") == ""


def test_vml_alt_text_is_leak_checked(tmp_path: Path) -> None:
    src = _docx_with(tmp_path / "in.docx", _VML)
    hidden = [ln for ln in Writer().extract_text(src).splitlines() if not ln.startswith("Dear ")]
    assert f"Photo of {NAME}" in hidden
    assert NAME in hidden  # v:imagedata o:title


def test_internal_link_tooltip_naming_a_redacted_value_is_blanked(tmp_path: Path) -> None:
    src = _docx_with(tmp_path / "in.docx", _ANCHOR_LINK)
    out = tmp_path / "out.docx"
    _engine().process_document(Reader(), Writer(), src, out, operator_policies=REPLACE)
    assert NAME.encode() not in _zip_bytes(out)
    [link] = docx.Document(str(out)).element.body.xpath(".//w:hyperlink")
    assert link.get(f"{{{_W}}}anchor") == "terms"  # the link itself still works


def test_unrelated_tooltip_and_alt_text_are_kept(tmp_path: Path) -> None:
    src = _docx_with(
        tmp_path / "in.docx",
        _ANCHOR_LINK.replace(f"Ask {NAME}", "Go to the terms"),
        _VML.replace(f"Photo of {NAME}", "Company logo").replace(
            f'o:title="{NAME}"', 'o:title="logo"'
        ),
    )
    out = tmp_path / "out.docx"
    _engine().process_document(Reader(), Writer(), src, out, operator_policies=REPLACE)
    raw = _zip_bytes(out)
    assert b"Go to the terms" in raw
    assert b"Company logo" in raw
