"""Names split across Word / PowerPoint runs are found and redacted whole.

Real Presidio stack. Every name below is invented.
"""

from __future__ import annotations

import json
import zipfile
from pathlib import Path

import docx
import pytest
from pptx import Presentation
from pptx.util import Inches
from sanctum.analyzer.adapter import PresidioAnalyzer
from sanctum.analyzer.recognizers import AnyDomainEmailRecognizer
from sanctum.anonymizer.adapter import PresidioAnonymizer
from sanctum.core.engine import SanctumEngine
from sanctum.core.exceptions import LeakCheckError
from sanctum.core.models import (
    DetectionResult,
    OperatorPolicy,
    ProposalDecision,
    UserAddedDecision,
)
from sanctum.core.review.session import add_decision, apply_user_added_with_overlap_purge
from sanctum.core.review.store import SessionStore
from sanctum.documents import pptx_adapter
from sanctum.documents.docx_adapter import Reader, Writer

pytestmark = pytest.mark.integration

SPLIT_RUNS = ["Dear Jennifer", " ", "Martin, thanks."]


@pytest.fixture(scope="module")
def engine() -> SanctumEngine:
    return SanctumEngine(analyzer=PresidioAnalyzer(), anonymizer=PresidioAnonymizer())


def make_docx(path: Path, runs: list[str]) -> Path:
    d = docx.Document()
    para = d.add_paragraph()
    for i, text in enumerate(runs):
        para.add_run(text).bold = i % 2 == 1  # alternate formatting forces separate runs
    d.save(str(path))
    return path


def make_pptx(path: Path, runs: list[str]) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    para = box.text_frame.paragraphs[0]
    for i, text in enumerate(runs):
        run = para.add_run()
        run.text = text
        run.font.bold = i % 2 == 1
    prs.save(str(path))
    return path


def accept_all(store: SessionStore, session_id: str) -> None:
    with store.locked(session_id):
        s = store.load(session_id)
        for p in s.proposals:
            add_decision(s, ProposalDecision(proposal_id=p.detection_id, status="accept"))
        store.save(s)


def test_split_name_is_redacted_whole_in_docx(engine: SanctumEngine, tmp_path: Path) -> None:
    store = SessionStore(root=tmp_path / "sessions")
    src = make_docx(tmp_path / "in.docx", SPLIT_RUNS)
    session = engine.create_review_session(
        Reader(), src, default_operator="replace", session_store=store
    )
    assert {p.group_original for p in session.proposals if p.group_id} == {"Jennifer Martin"}
    accept_all(store, session.id)
    out = tmp_path / "out.docx"
    engine.commit_review_session(Reader(), Writer(), session.id, out, store)
    text = "".join(r.text for r in docx.Document(str(out)).paragraphs[0].runs)
    assert text == "Dear <PERSON>, thanks."


def test_rejected_split_name_survives_and_passes_the_leak_check(
    engine: SanctumEngine, tmp_path: Path
) -> None:
    store = SessionStore(root=tmp_path / "sessions")
    src = make_docx(tmp_path / "in.docx", SPLIT_RUNS)
    session = engine.create_review_session(
        Reader(), src, default_operator="replace", session_store=store
    )
    tail = next(p for p in session.proposals if p.group_id and p.group_index > 0)
    with store.locked(session.id):
        s = store.load(session.id)
        add_decision(s, ProposalDecision(proposal_id=tail.detection_id, status="reject"))
        store.save(s)
    out = tmp_path / "out.docx"
    engine.commit_review_session(Reader(), Writer(), session.id, out, store)
    assert docx.Document(str(out)).paragraphs[0].text == "".join(SPLIT_RUNS)


def test_split_name_is_redacted_whole_in_docx_without_review(
    engine: SanctumEngine, tmp_path: Path
) -> None:
    src = make_docx(tmp_path / "in.docx", SPLIT_RUNS)
    out = tmp_path / "out.docx"
    replace = {"DEFAULT": OperatorPolicy(operator_name="replace")}
    engine.process_document(Reader(), Writer(), src, out, operator_policies=replace)
    assert docx.Document(str(out)).paragraphs[0].text == "Dear <PERSON>, thanks."


def test_split_name_is_redacted_whole_in_pptx(engine: SanctumEngine, tmp_path: Path) -> None:
    store = SessionStore(root=tmp_path / "sessions")
    src = make_pptx(tmp_path / "in.pptx", SPLIT_RUNS)
    reader, writer = pptx_adapter.Reader(), pptx_adapter.Writer()
    session = engine.create_review_session(
        reader, src, default_operator="replace", session_store=store
    )
    assert {p.group_original for p in session.proposals if p.group_id} == {"Jennifer Martin"}
    accept_all(store, session.id)
    out = tmp_path / "out.pptx"
    engine.commit_review_session(reader, writer, session.id, out, store)
    shape = Presentation(str(out)).slides[0].shapes[0]
    assert shape.text_frame.text == "Dear <PERSON>, thanks."


def test_rc3_manifest_without_new_fields_still_commits(
    engine: SanctumEngine, tmp_path: Path
) -> None:
    store = SessionStore(root=tmp_path / "sessions")
    src = make_docx(tmp_path / "in.docx", ["Dear Cameron Dean, thanks."])
    session = engine.create_review_session(
        Reader(), src, default_operator="replace", session_store=store
    )
    manifest = store.root / session.id / "manifest.json"
    raw = json.loads(manifest.read_text())
    for seg in raw["segments"]:
        seg.pop("block", None)
        seg.pop("join_before", None)
    for prop in raw["proposals"]:
        for key in ("group_id", "group_index", "group_original"):
            prop.pop(key, None)
    manifest.write_text(json.dumps(raw))

    accept_all(store, session.id)
    out = tmp_path / "out.docx"
    engine.commit_review_session(Reader(), Writer(), session.id, out, store)
    assert "Cameron" not in docx.Document(str(out)).paragraphs[0].text


def test_split_internal_domain_email_is_one_email_finding(tmp_path: Path) -> None:
    # Production wiring: the any-domain email recognizer is registered (see _create_engine).
    engine = SanctumEngine(
        analyzer=PresidioAnalyzer(extra_recognizers=[AnyDomainEmailRecognizer()]),
        anonymizer=PresidioAnonymizer(),
    )
    store = SessionStore(root=tmp_path / "sessions")
    src = make_docx(tmp_path / "in.docx", ["Contact j.al", "brecht@northgate.local today"])
    session = engine.create_review_session(
        Reader(), src, default_operator="replace", session_store=store
    )
    emails = [p for p in session.proposals if p.entity_type == "EMAIL_ADDRESS"]
    assert emails
    assert {p.group_original for p in emails} == {"j.albrecht@northgate.local"}
    assert not [p for p in session.proposals if p.entity_type == "URL"]


def test_committed_docx_has_no_author_or_comments(engine: SanctumEngine, tmp_path: Path) -> None:
    store = SessionStore(root=tmp_path / "sessions")
    d = docx.Document()
    d.core_properties.author = "Jennifer Martin"
    d.core_properties.last_modified_by = "Jennifer Martin"
    para = d.add_paragraph()
    run = para.add_run("The contract was signed on Monday.")
    d.add_comment(run, text="Ping Jennifer Martin", author="Jennifer Martin", initials="JM")
    src = tmp_path / "in.docx"
    d.save(str(src))
    session = engine.create_review_session(
        Reader(), src, default_operator="replace", session_store=store
    )
    accept_all(store, session.id)
    out = tmp_path / "out.docx"
    engine.commit_review_session(Reader(), Writer(), session.id, out, store)
    with zipfile.ZipFile(out) as z:
        assert not any(n.startswith("word/comments") for n in z.namelist())
        assert b"Jennifer Martin" not in b"".join(z.read(n) for n in z.namelist())


def test_committed_pptx_has_no_author(engine: SanctumEngine, tmp_path: Path) -> None:
    store = SessionStore(root=tmp_path / "sessions")
    src = make_pptx(tmp_path / "in.pptx", SPLIT_RUNS)
    prs = Presentation(str(src))
    prs.core_properties.author = "Jennifer Martin"
    prs.save(str(src))
    reader, writer = pptx_adapter.Reader(), pptx_adapter.Writer()
    session = engine.create_review_session(
        reader, src, default_operator="replace", session_store=store
    )
    accept_all(store, session.id)
    out = tmp_path / "out.pptx"
    engine.commit_review_session(reader, writer, session.id, out, store)
    with zipfile.ZipFile(out) as z:
        assert b"Jennifer Martin" not in b"".join(z.read(n) for n in z.namelist())


def _docx_with_header(path: Path, header_runs: list[str]) -> Path:
    d = docx.Document()
    para = d.add_paragraph()
    for i, text in enumerate(SPLIT_RUNS):
        para.add_run(text).bold = i % 2 == 1
    header = d.sections[0].header.paragraphs[0]
    for i, text in enumerate(header_runs):
        header.add_run(text).bold = i % 2 == 0
    d.save(str(path))
    return path


def test_name_in_a_docx_header_is_redacted(engine: SanctumEngine, tmp_path: Path) -> None:
    # Header runs are segments like body runs: detected (even split across
    # runs), redacted, and written back, so the leak check passes.
    src = _docx_with_header(tmp_path / "in.docx", ["Prepared for Jen", "nifer Martin"])
    out = tmp_path / "out.docx"
    replace = {"DEFAULT": OperatorPolicy(operator_name="replace")}
    engine.process_document(Reader(), Writer(), src, out, operator_policies=replace)
    written = docx.Document(str(out))
    assert written.paragraphs[0].text == "Dear <PERSON>, thanks."
    assert written.sections[0].header.paragraphs[0].text == "Prepared for <PERSON>"


class _BodyOnlyAnalyzer:
    """Finds "Jennifer Martin" only in text that starts with "Dear" (i.e. not the header)."""

    def analyze(
        self,
        text: str,
        language: str = "en",
        entities: list[str] | None = None,
        score_threshold: float | None = None,
    ) -> list[DetectionResult]:
        i = text.find("Jennifer Martin")
        if not text.startswith("Dear") or i < 0:
            return []
        return [
            DetectionResult(
                entity_type="PERSON", start=i, end=i + 15, score=0.9, text_span="Jennifer Martin"
            )
        ]


def test_name_missed_in_a_docx_header_fails_closed(tmp_path: Path) -> None:
    # If detection misses the header copy of a name it replaced in the body,
    # the leak check still sees the header and refuses the output.
    engine = SanctumEngine(analyzer=_BodyOnlyAnalyzer(), anonymizer=PresidioAnonymizer())
    src = _docx_with_header(tmp_path / "in.docx", ["Prepared for Jennifer Martin"])
    out = tmp_path / "out.docx"
    replace = {"DEFAULT": OperatorPolicy(operator_name="replace")}
    with pytest.raises(LeakCheckError) as exc_info:
        engine.process_document(Reader(), Writer(), src, out, operator_policies=replace)
    assert "Jennifer Martin" in exc_info.value.leaks
    assert not out.exists()


def test_hand_marking_part_of_a_linked_name_keeps_the_rest_redacted(
    engine: SanctumEngine, tmp_path: Path
) -> None:
    # C1 / Ruling 12: marking "Jennifer" by hand must not un-redact "Martin".
    store = SessionStore(root=tmp_path / "sessions")
    src = make_docx(tmp_path / "in.docx", ["Dear Jennifer", " Martin, thanks."])
    session = engine.create_review_session(
        Reader(), src, default_operator="replace", session_store=store
    )
    assert {p.group_original for p in session.proposals if p.group_id} == {"Jennifer Martin"}
    accept_all(store, session.id)
    with store.locked(session.id):
        s = store.load(session.id)
        ua = UserAddedDecision(
            segment_anchor="body/p0/r0", entity_type="PERSON", original="Jennifer", start=5, end=13
        )
        apply_user_added_with_overlap_purge(s, ua)
        store.save(s)
        assert [p.original for p in s.proposals] == ["Martin"]
    out = tmp_path / "out.docx"
    engine.commit_review_session(Reader(), Writer(), session.id, out, store)
    text = docx.Document(str(out)).paragraphs[0].text
    assert "Martin" not in text and "Jennifer" not in text


def _add_user_added(store: SessionStore, session_id: str, ua: UserAddedDecision) -> None:
    with store.locked(session_id):
        s = store.load(session_id)
        apply_user_added_with_overlap_purge(s, ua)
        store.save(s)


def test_user_added_span_lands_where_it_was_marked(engine: SanctumEngine, tmp_path: Path) -> None:
    # I1: the second "Martin" is hand-marked; the span must land at offset 21,
    # not on the first occurrence (which an accepted proposal already covers).
    store = SessionStore(root=tmp_path / "sessions")
    text = "Martin called. Later Martin left."
    src = make_docx(tmp_path / "in.docx", [text])
    session = engine.create_review_session(
        Reader(), src, default_operator="replace", session_store=store
    )
    with store.locked(session.id):
        s = store.load(session.id)
        # Keep only the first "Martin" as a detected proposal, as in the review.
        s.proposals = [p for p in s.proposals if p.original == "Martin" and p.start == 0]
        assert len(s.proposals) == 1
        add_decision(s, ProposalDecision(proposal_id=s.proposals[0].detection_id, status="accept"))
        store.save(s)
    _add_user_added(
        store,
        session.id,
        UserAddedDecision(
            segment_anchor="body/p0/r0",
            entity_type="USER_ADDED",
            original="Martin",
            start=21,
            end=27,
        ),
    )
    out = tmp_path / "out.docx"
    engine.commit_review_session(Reader(), Writer(), session.id, out, store)
    assert docx.Document(str(out)).paragraphs[0].text == "<PERSON> called. Later [REDACTED] left."


def test_user_added_span_commits_as_redacted_under_default_params(
    engine: SanctumEngine, tmp_path: Path
) -> None:
    store = SessionStore(root=tmp_path / "sessions")
    src = make_docx(tmp_path / "in.docx", ["Ask the plumber about code 4471 today."])
    session = engine.create_review_session(
        Reader(), src, default_operator="replace", session_store=store
    )
    _add_user_added(
        store,
        session.id,
        UserAddedDecision(
            segment_anchor="body/p0/r0",
            entity_type="USER_ADDED",
            original="plumber",
            start=8,
            end=15,
        ),
    )
    out = tmp_path / "out.docx"
    engine.commit_review_session(Reader(), Writer(), session.id, out, store)
    assert "Ask the [REDACTED] about" in docx.Document(str(out)).paragraphs[0].text
